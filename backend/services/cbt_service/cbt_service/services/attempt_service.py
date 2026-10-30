import asyncio
from dataclasses import dataclass
from uuid import UUID, uuid4
from datetime import datetime, timezone, timedelta
from fastapi import HTTPException
from sqlmodel import Session, delete, select, col, delete, func, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError

from pydantic import BaseModel
from typing import Any
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
from fastapi.responses import StreamingResponse


from cbt_service.database.models.violation import ExamViolation
from cbt_service.database.models.enums.exam_enum import AssignmentStatus, ExamAction, ExamStatus
from cbt_service.schemas.item_subject_schemas import ItemForDisplayRead, ItemForScoringRead
from cbt_service.services.item.item_service import ItemService
from cbt_service.database.models.exam import ExamAssignment, ExamItem, ExamModel
from cbt_service.database.models.item import ItemModel
from cbt_service.services.marking.scoring import pass_threshold, score_answer, verdict
from cbt_service.services.exam_service import ExamService, _log
from cbt_service.services.paging import page_of, total_pages
from sqlmodel.sql.expression import SelectOfScalar
from cbt_service.database.models.attempt import AttemptModel, LateResponseModel, ResponseModel

from cbt_service.schemas.attempt_schemas import (
    AttemptDetailRead,
    AttemptExamRead,
    AttemptItemRead,
    AttemptResponseDetail,
    AttemptSectionRead,
    CohortAttemptSummary,
    ResetAttemptRequest,
    StartAttemptRequest,
    SaveResponseRequest,
    ManualReviewRequest,
    AttemptRead,
    AttemptStatusRead,
    AttemptReviewRead,
    AttemptOptionRead,
    LateAnswerRead,
    AttemptPage,
    ResponseRead,
    ReviewItem,
    CurrentUser,
)
from cbt_service.database.models.enums.attempt_enum import AttemptStatus
from cbt_service.core.permissions import P


@dataclass
class ExamExportData:
    exam_items: list[Any]
    attempts: list[AttemptModel]
    responses: list[ResponseModel]

# ============================================================
# SERVICE
# ============================================================

# Question types whose options keep their order (shuffling True/False reads oddly).
FIXED_OPTION_TYPES: frozenset[str] = frozenset({"true_false"})
# Time is strict: an attempt closes at its deadline and nothing saved after it
# changes the score. Answers held up by a bad connection that arrive within
# LATE_WINDOW are kept as late answers for staff to accept or reject, with
# evidence (how late, when the last on-time save was).
LATE_WINDOW: timedelta = timedelta(minutes=10)
# The browser's auto-submit at 0:00 can land a moment before the server's
# deadline; such an attempt still counts as closed by time.
CLOCK_SLACK: timedelta = timedelta(seconds=10)

CLOSED_ATTEMPT_STATUSES: tuple[AttemptStatus, ...] = (
    AttemptStatus.SUBMITTED, AttemptStatus.SCORED, AttemptStatus.TERMINATED,
)


class AttemptService:

    @staticmethod
    def _set_assignment_status(session: Session, attempt: AttemptModel, status: AssignmentStatus) -> None:
        """Keep the assignment's status in step with the attempt (counts and lists use it)."""
        assignment: ExamAssignment | None = (
            session.get(ExamAssignment, attempt.assignment_id) if attempt.assignment_id else None
        )
        if assignment is not None and assignment.status != status:
            assignment.status = status
            session.add(assignment)

    @staticmethod
    def _aware(dt: datetime) -> datetime:
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)

    @staticmethod
    def start(session: Session, payload: StartAttemptRequest, student_id: UUID, org_id: UUID) -> AttemptModel:
        exam = session.exec(
            select(ExamModel).where(
                ExamModel.id == payload.exam_id,
                ExamModel.org_id == org_id,
            )
        ).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Exam not found.")

        # assignment = session.exec(
        #     select(ExamAssignment).where(ExamAssignment.id == payload.assignment_id)
        # ).first()
        # Lock the assignment row: two "start" requests arriving together (a
        # double click, two tabs, a retry on a slow connection) wait for each
        # other here, so the second sees the attempt the first created instead
        # of creating a duplicate.
        assignment = session.exec(
            select(ExamAssignment).where(
                ExamAssignment.id == payload.assignment_id,
                ExamAssignment.exam_id == payload.exam_id,
                ExamAssignment.student_id == student_id,
                ExamAssignment.org_id == org_id,
            ).with_for_update()
        ).first()
        if not assignment:
            raise HTTPException(status_code=404, detail="Assignment not found.")

        existing = session.exec(
            select(AttemptModel).where(
                AttemptModel.exam_id == payload.exam_id,
                AttemptModel.student_id == student_id,
            )
        ).all()

        effective_start, effective_end = ExamService.get_effective_window(exam, assignment)
        now = datetime.now(timezone.utc)

        if effective_start and now < effective_start:
            raise HTTPException(status_code=400, detail="This exam is not yet available.")
        if effective_end and now > effective_end:
            raise HTTPException(status_code=400, detail="The window for this exam has closed.")

        in_progress = next((a for a in existing if a.status == AttemptStatus.STARTED), None)
        if in_progress:
            return in_progress

        # Reactivate a reset attempt in place, instead of creating a new row.
        # A reset is meant to give the SAME attempt a clean restart, not free
        # up a slot for a separate attempt record — otherwise you end up with
        # one permanently-RESET row and a second independent SUBMITTED row.
        reset_attempts = [a for a in existing if a.status == AttemptStatus.RESET]
        if reset_attempts:
            # If more than one has ever been reset, resume the most recent.
            attempt_to_resume = max(reset_attempts, key=lambda a: a.started_at)
            attempt_to_resume.status = AttemptStatus.STARTED
            attempt_to_resume.started_at = datetime.now(timezone.utc)
            attempt_to_resume.clock_started_at = None      # starts when the questions are sent
            attempt_to_resume.submitted_at = None
            attempt_to_resume.raw_score = None
            attempt_to_resume.final_score = None
            attempt_to_resume.percentage = None
            attempt_to_resume.passed = None
            attempt_to_resume.scored_at = None
            attempt_to_resume.scored_by = None
            session.add(attempt_to_resume)
            AttemptService._set_assignment_status(session, attempt_to_resume, AssignmentStatus.STARTED)
            session.commit()
            session.refresh(attempt_to_resume)
            return attempt_to_resume

        max_attempts = exam.max_attempts or 1
        countable = [a for a in existing if a.status != AttemptStatus.RESET]
        if len(countable) >= max_attempts:
            raise HTTPException(
                status_code=400,
                detail=f"Maximum attempts ({max_attempts}) reached for this exam.",
            )

        attempt = AttemptModel(
            exam_id=payload.exam_id,
            student_id=student_id,
            org_id=org_id,
            assignment_id=payload.assignment_id,
            attempt_number=max((a.attempt_number for a in existing), default=0) + 1,
        )
        session.add(attempt)
        AttemptService._set_assignment_status(session, attempt, AssignmentStatus.STARTED)
        try:
            session.commit()
        except IntegrityError:
            # The database's own guard (one row per attempt number): another
            # request created this attempt a moment ago. Return that one.
            session.rollback()
            same = session.exec(
                select(AttemptModel).where(
                    AttemptModel.exam_id == payload.exam_id,
                    AttemptModel.student_id == student_id,
                    AttemptModel.attempt_number == attempt.attempt_number,
                )
            ).first()
            if same is None:
                raise
            return same
        session.refresh(attempt)
        return attempt


    # @staticmethod
    # def save_response_new(
    #     session: Session, attempt_id: UUID, payload: SaveResponseRequest, student_id: UUID
    # ) -> ResponseModel:
    #     attempt = AttemptService.get_attempt(session, attempt_id, student_id)
    #     if attempt.status != AttemptStatus.STARTED:
    #         raise HTTPException(
    #             status_code=409,
    #             detail={"code": "ATTEMPT_CLOSED", "status": attempt.status.value},
    #         )

    #     now = datetime.now(timezone.utc)
    #     stmt = (
    #         pg_insert(ResponseModel)
    #         .values(
    #             id=uuid4(),  # SQLModel default_factory is NOT applied on Core inserts
    #             attempt_id=attempt_id,
    #             item_id=payload.item_id,
    #             exam_item_id=payload.exam_item_id,
    #             org_id=attempt.org_id,
    #             answer=payload.answer,
    #             time_spent_seconds=payload.time_spent_seconds,
    #             is_flagged=payload.is_flagged,
    #             updated_at=now,
    #         )
    #         .on_conflict_do_update(
    #             index_elements=["attempt_id", "item_id"],
    #             set_={
    #                 "answer": payload.answer,
    #                 "time_spent_seconds": payload.time_spent_seconds,
    #                 "is_flagged": payload.is_flagged,
    #                 "updated_at": now,
    #             },
    #         )
    #         .returning(ResponseModel)
    #     )
    #     response: ResponseModel = session.exec(stmt).scalar_one()
    #     session.commit()
    #     return response

    @staticmethod
    def _allowed(exam: ExamModel) -> timedelta:
        """The candidate's full time: the exam's duration plus its extra time."""
        return timedelta(minutes=exam.duration_minutes, seconds=exam.extra_time_seconds or 0)

    @staticmethod
    def _deadline(attempt: AttemptModel, exam: ExamModel) -> datetime | None:
        """When time runs out; None while the clock hasn't started (questions not yet sent)."""
        if attempt.clock_started_at is None:
            return None
        return AttemptService._aware(attempt.clock_started_at) + AttemptService._allowed(exam)

    @staticmethod
    def _timing(attempt: AttemptModel, exam: ExamModel | None, now: datetime) -> tuple[datetime, int]:
        """(expires_at, remaining seconds) for the timer. Before the clock starts: the full time."""
        if exam is None:
            return now, 0
        deadline = AttemptService._deadline(attempt, exam)
        expires_at = deadline if deadline is not None else now + AttemptService._allowed(exam)
        remaining = (
            max(0, int((expires_at - now).total_seconds())) if attempt.status == AttemptStatus.STARTED else 0
        )
        return expires_at, remaining

    @staticmethod
    def save_response(
        session: Session,
        attempt_id: UUID,
        payload: SaveResponseRequest,
        student_id: UUID,
    ) -> ResponseModel | LateResponseModel:
        """
        Save an answer while the attempt is open. After time runs out, an
        answer that arrives within LATE_WINDOW (held up by the connection) is
        kept as a late answer for staff to decide on; it doesn't change the score.
        """
        attempt = session.get(AttemptModel, attempt_id)
        if attempt is None or attempt.student_id != student_id:
            raise HTTPException(status_code=404, detail="Attempt not found.")
        exam = session.get(ExamModel, attempt.exam_id)
        if exam is None:
            raise HTTPException(status_code=404, detail="Exam not found.")

        # The question must be on this exam; its exam entry comes from here,
        # not from the client.
        exam_item: ExamItem | None = session.exec(
            select(ExamItem).where(ExamItem.exam_id == attempt.exam_id, ExamItem.item_id == payload.item_id)
        ).first()

        now = datetime.now(timezone.utc)
        deadline = AttemptService._deadline(attempt, exam)
        if attempt.status == AttemptStatus.STARTED and deadline is not None and now >= deadline:
            attempt = AttemptService._finalize_submission(session, attempt, exam, auto=True)

        if attempt.status != AttemptStatus.STARTED:
            closed_by_time: bool = (
                deadline is not None
                and attempt.status in (AttemptStatus.SUBMITTED, AttemptStatus.SCORED)
                and attempt.submitted_at is not None
                and AttemptService._aware(attempt.submitted_at) >= deadline - CLOCK_SLACK
            )
            if exam_item is not None and closed_by_time and deadline is not None and now <= deadline + LATE_WINDOW:
                return AttemptService._keep_late_answer(session, attempt, exam_item, payload, now, deadline)
            raise HTTPException(
                status_code=409,
                detail={"code": "ATTEMPT_CLOSED", "status": attempt.status.value},
            )

        if exam_item is None:
            raise HTTPException(status_code=422, detail="This question isn't part of the exam.")

        existing = session.exec(
            select(ResponseModel).where(
                ResponseModel.attempt_id == attempt_id,
                ResponseModel.item_id == payload.item_id,
            )
        ).first()

        if existing:
            existing.answer = payload.answer
            existing.time_spent_seconds = payload.time_spent_seconds
            existing.is_flagged = payload.is_flagged
            existing.updated_at = now
            session.add(existing)
            session.commit()
            session.refresh(existing)
            return existing

        response = ResponseModel(
            attempt_id=attempt_id,
            item_id=payload.item_id,
            exam_item_id=exam_item.id,
            org_id=attempt.org_id,
            answer=payload.answer,
            time_spent_seconds=payload.time_spent_seconds,
            is_flagged=payload.is_flagged,
        )
        session.add(response)
        session.commit()
        session.refresh(response)
        return response

    @staticmethod
    def _keep_late_answer(
        session: Session, attempt: AttemptModel, exam_item: ExamItem,
        payload: SaveResponseRequest, now: datetime, deadline: datetime,
    ) -> LateResponseModel:
        """One late answer per question: a newer one replaces an undecided older one."""
        last_on_time: datetime | None = session.exec(
            select(func.max(ResponseModel.updated_at)).where(
                ResponseModel.attempt_id == attempt.id, col(ResponseModel.updated_at) <= deadline,
            )
        ).one()
        late = session.exec(
            select(LateResponseModel).where(
                LateResponseModel.attempt_id == attempt.id,
                LateResponseModel.item_id == exam_item.item_id,
                LateResponseModel.status == "pending",
            )
        ).first() or LateResponseModel(
            attempt_id=attempt.id, org_id=attempt.org_id,
            item_id=exam_item.item_id, exam_item_id=exam_item.id,
        )
        late.answer = payload.answer
        late.received_at = now
        late.seconds_late = round((now - deadline).total_seconds(), 1)
        late.last_on_time_save_at = last_on_time
        session.add(late)
        session.commit()
        session.refresh(late)
        return late

    @staticmethod
    def late_pending_counts(session: Session, attempt_ids: list[UUID]) -> dict[UUID, int]:
        """Per attempt: late answers waiting for a decision."""
        if not attempt_ids:
            return {}
        rows = session.exec(
            select(LateResponseModel.attempt_id, func.count())
            .where(col(LateResponseModel.attempt_id).in_(attempt_ids), LateResponseModel.status == "pending")
            .group_by(LateResponseModel.attempt_id)
        ).all()
        return {attempt_id: int(n) for attempt_id, n in rows}

    @staticmethod
    def late_answers_for(
        session: Session, attempt: AttemptModel, items: dict[UUID, ItemModel], responses: list[ResponseModel],
    ) -> list[LateAnswerRead]:
        lates = session.exec(
            select(LateResponseModel).where(LateResponseModel.attempt_id == attempt.id)
            .order_by(col(LateResponseModel.received_at))
        ).all()
        if not lates:
            return []
        missing = [l.item_id for l in lates if l.item_id not in items]
        if missing:
            items = {**items, **{i.id: i for i in session.exec(select(ItemModel).where(col(ItemModel.id).in_(missing))).all()}}
        saved: dict[UUID, ResponseModel] = {r.item_id: r for r in responses}
        out: list[LateAnswerRead] = []
        for l in lates:
            item = items.get(l.item_id)
            current = saved.get(l.item_id)
            out.append(LateAnswerRead(
                id=l.id, item_id=l.item_id, exam_item_id=l.exam_item_id,
                question_text=item.question_text if item else "[question no longer available]",
                item_type=str(getattr(item.item_type, "value", item.item_type)) if item else None,
                answer=l.answer, current_answer=current.answer if current else None,
                received_at=l.received_at, seconds_late=l.seconds_late,
                last_on_time_save_at=l.last_on_time_save_at, status=l.status, decided_at=l.decided_at,
            ))
        return out

    @staticmethod
    def decide_late_answer(
        session: Session, late_id: UUID, accept: bool, current_user: CurrentUser,
    ) -> AttemptDetailRead:
        """
        Accept: the late answer replaces the saved one and is marked like any
        other (auto-marked, or sent to the marking queue); the attempt is re-scored.
        Reject: kept on record, the score is unchanged.
        """
        late = session.get(LateResponseModel, late_id)
        if late is None or late.org_id != current_user.org_id:
            raise HTTPException(status_code=404, detail="Late answer not found.")
        attempt = session.get(AttemptModel, late.attempt_id)
        exam = session.get(ExamModel, attempt.exam_id) if attempt else None
        if attempt is None or exam is None:
            raise HTTPException(status_code=404, detail="Attempt not found.")
        ExamService._assert_can_view(session, exam, current_user, P.RESULTS_MARK)
        if late.status != "pending":
            raise HTTPException(status_code=409, detail=f"This late answer was already {late.status}.")

        if accept:
            exam_item = session.exec(
                select(ExamItem).where(ExamItem.exam_id == exam.id, ExamItem.item_id == late.item_id)
            ).first()
            if exam_item is None:
                raise HTTPException(status_code=409, detail="This question is no longer on the exam.")
            item_data = AttemptService._get_items_for_scoring(session, [late.item_id]).get(late.item_id)
            response = session.exec(
                select(ResponseModel).where(
                    ResponseModel.attempt_id == attempt.id, ResponseModel.item_id == late.item_id,
                )
            ).first() or ResponseModel(
                attempt_id=attempt.id, item_id=late.item_id, exam_item_id=exam_item.id, org_id=attempt.org_id,
            )
            response.answer = late.answer
            response.updated_at = datetime.now(timezone.utc)
            response.reviewed_by = None
            response.reviewed_at = None
            if item_data:
                score = score_answer(
                    str(getattr(item_data.get("item_type"), "value", item_data.get("item_type"))),
                    item_data.get("correct_answers"), late.answer, exam_item.marks, exam_item.negative_marks,
                    auto_mark_short_answers=bool(exam.auto_mark_short_answers),
                )
                response.is_correct = score.is_correct
                response.marks_awarded = score.marks_awarded
            session.add(response)
            session.flush()
            responses = list(session.exec(select(ResponseModel).where(ResponseModel.attempt_id == attempt.id)).all())
            AttemptService.recompute_score(session, attempt, exam, responses)
            session.add(attempt)

        late.status = "accepted" if accept else "rejected"
        late.decided_by = current_user.id
        late.decided_at = datetime.now(timezone.utc)
        session.add(late)
        _log(session, exam.id, exam.org_id, current_user.id, ExamAction.UPDATED, {
            "field": "late_answer", "attempt_id": str(attempt.id), "item_id": str(late.item_id),
            "decision": late.status, "seconds_late": late.seconds_late,
        })
        session.commit()
        return AttemptService.get_attempt_detail_for_staff(session, attempt.id, current_user)

    @staticmethod
    def submit(
        session: Session,
        attempt_id: UUID,
        student_id: UUID,
        pass_mark: float | None = None,
    ) -> AttemptModel:
        attempt = AttemptService.get_attempt(session, attempt_id, student_id)  # expiry-checked

        if attempt.status != AttemptStatus.STARTED:
            return attempt
            # raise HTTPException(status_code=400, detail="Attempt already submitted.")

        exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Exam not found.")

        return AttemptService._finalize_submission(session, attempt, exam, pass_mark=pass_mark, auto=False)
    

    @staticmethod
    def manual_review(
        session: Session,
        response_id: UUID,
        payload: ManualReviewRequest,
        reviewer_id: UUID,
        current_user: CurrentUser | None = None,
    ) -> ResponseModel:
        resp = session.exec(
            select(ResponseModel).where(ResponseModel.id == response_id)
        ).first()

        if not resp:
            raise HTTPException(status_code=404, detail="Response not found.")

        if current_user is not None:
            # Only answers to exams the marker can reach, in their own organisation.
            owner = session.exec(
                select(AttemptModel).where(
                    AttemptModel.id == resp.attempt_id,
                    AttemptModel.org_id == current_user.org_id,
                )
            ).first()
            if owner is None:
                raise HTTPException(status_code=404, detail="Response not found.")
            exam_for_check: ExamModel = ExamService._get_exam(session, owner.exam_id, current_user.org_id)
            ExamService._assert_can_view(session, exam_for_check, current_user, P.RESULTS_MARK)

        attempt = session.exec(
            select(AttemptModel).where(AttemptModel.id == resp.attempt_id)
        ).first()
        if attempt is None:
            raise HTTPException(status_code=404, detail="Response not found.")
        if attempt.status in (AttemptStatus.STARTED, AttemptStatus.RESET):
            raise HTTPException(status_code=409, detail="This attempt hasn't been submitted yet.")
        exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
        if exam is None:
            raise HTTPException(status_code=404, detail="Exam not found.")

        exam_item = AttemptService._exam_items(session, exam.id).get(resp.item_id)
        if exam_item is None:
            raise HTTPException(status_code=422, detail="This answer isn't for a question on the exam.")
        max_marks: float = exam_item.marks
        if payload.marks_awarded < 0 or payload.marks_awarded > max_marks:
            raise HTTPException(
                status_code=422,
                detail=f"Marks must be between 0 and {max_marks:g} for this question.",
            )

        resp.marks_awarded = payload.marks_awarded
        resp.is_correct = verdict(payload.marks_awarded, max_marks)
        resp.reviewed_by = reviewer_id
        resp.review_notes = (payload.review_notes or "").strip() or None
        resp.reviewed_at = datetime.now(timezone.utc)
        session.add(resp)
        session.flush()

        attempt.scored_by = str(reviewer_id)
        AttemptService.recompute_score(session, attempt, exam)

        session.commit()
        session.refresh(resp)
        return resp

    @staticmethod
    def get_attempt(session: Session, attempt_id: UUID, user_id: UUID) -> AttemptModel:
        attempt = session.exec(
            select(AttemptModel).where(AttemptModel.id == attempt_id)
        ).first()
        if not attempt or attempt.student_id != user_id:
            raise HTTPException(status_code=404, detail="Attempt not found.")
        return AttemptService._enforce_time_limit(session, attempt)
        # return attempt

    @staticmethod
    def get_attempt_status(
        session: Session, attempt_id: UUID, student_id: UUID
    ) -> AttemptStatusRead:
        """
        Heartbeat poll. Loads the attempt, its exam and the live violation
        count in ONE query (previously 5 round trips, 3 of them for the same
        exam row). Still enforces expiry: an attempt past its deadline is
        auto-submitted, exactly as get_attempt() does.
        """
        # Violations count for this attempt, computed inside the same query.
        # (Inline rather than ViolationService._count_violations to avoid a
        # circular import: violation_service imports this module.)
        violation_count = (
            select(func.count())
            .select_from(ExamViolation)
            .where(
                ExamViolation.attempt_id == AttemptModel.id,
                col(ExamViolation.cleared_at).is_(None),
            )
            .correlate(AttemptModel)
            .scalar_subquery()
        )

        row = session.exec(
            select(AttemptModel, ExamModel, violation_count)
            .outerjoin(ExamModel, col(ExamModel.id) == AttemptModel.exam_id)
            .where(
                AttemptModel.id == attempt_id,
                AttemptModel.student_id == student_id,  # ownership: 404 otherwise
            )
        ).first()
        if row is None:
            raise HTTPException(status_code=404, detail="Attempt not found.")

        attempt: AttemptModel = row[0]
        exam: ExamModel | None = row[1]
        count: int = row[2]

        if exam is not None:
            # Only writes (auto-submit) when the deadline has passed.
            attempt = AttemptService._enforce_time_limit_new(session, attempt, exam)

        read: AttemptRead = AttemptService._to_attempt_read_new(attempt, exam)
        return AttemptStatusRead(
            status=attempt.status,
            remaining_seconds=read.remaining_seconds,
            violation_count=count,
            max_violations=exam.max_violations if exam else None,
            server_time=read.server_time,
            clock_started=read.clock_started,
        )

    @staticmethod
    def get_responses(session: Session, attempt_id: UUID, user_id: UUID) -> list[ResponseModel]:
        AttemptService.get_attempt(session, attempt_id, user_id)  # expiry-checked
        return session.exec(
            select(ResponseModel).where(ResponseModel.attempt_id == attempt_id)
        ).all()
    


    @staticmethod
    def responses_for_student(session: Session, attempt_id: UUID, user_id: UUID) -> list[ResponseRead]:
        """The candidate's saved answers; marks left out unless the exam shows results."""
        attempt = AttemptService.get_attempt(session, attempt_id, user_id)
        exam = session.get(ExamModel, attempt.exam_id)
        hide: bool = exam is not None and not exam.show_result_immediately
        out: list[ResponseRead] = []
        for r in session.exec(select(ResponseModel).where(ResponseModel.attempt_id == attempt_id)).all():
            read = ResponseRead.model_validate(r, from_attributes=True)
            if hide:
                read.is_correct = read.marks_awarded = read.review_notes = None
            out.append(read)
        return out

    @staticmethod
    def review_for_student(session: Session, attempt_id: UUID, user_id: UUID) -> AttemptReviewRead:
        """
        A candidate's own look back at a finished attempt: only when the exam
        allows review. Marks, correct answers and comments only when it also
        shows results. Marking guides are never included.
        """
        attempt = AttemptService.get_attempt(session, attempt_id, user_id)
        if attempt.status not in CLOSED_ATTEMPT_STATUSES:
            raise HTTPException(status_code=409, detail="You can review this exam once it's submitted.")
        exam = session.get(ExamModel, attempt.exam_id)
        if exam is None:
            raise HTTPException(status_code=404, detail="Exam not found.")
        if not exam.allow_review:
            raise HTTPException(status_code=403, detail="Reviewing answers isn't allowed for this exam.")
        show: bool = bool(exam.show_result_immediately)

        sections = ExamService.get_sections_with_items_internal(session, exam.id)
        exam_items = [ei for section in sections for ei in section.items]
        items: dict[UUID, ItemModel] = {
            i.id: i for i in session.exec(
                select(ItemModel).where(col(ItemModel.id).in_([ei.item_id for ei in exam_items]))
            ).all()
        } if exam_items else {}
        responses: dict[UUID, ResponseModel] = {
            r.item_id: r for r in session.exec(
                select(ResponseModel).where(ResponseModel.attempt_id == attempt.id)
            ).all()
        }
        out: list[ReviewItem] = []
        for ei in exam_items:
            item = items.get(ei.item_id)
            if item is None:
                continue
            item_type: str = str(getattr(item.item_type, "value", item.item_type))
            r = responses.get(ei.item_id)
            written: bool = item_type in ("short_answer", "essay")
            options = [
                AttemptOptionRead(key=str(o.get("key")), text=str(o.get("text", "")))
                for o in (item.options or []) if isinstance(o, dict)
            ] or None
            out.append(ReviewItem(
                item_id=item.id, question_text=item.question_text, item_type=item_type,
                options=options, answer=r.answer if r is not None else None, marks=ei.marks,
                marks_awarded=r.marks_awarded if show and r is not None else None,
                is_correct=r.is_correct if show and r is not None else None,
                correct_answers=(
                    item.correct_answers if show and item_type != "essay" and item.correct_answers else None
                ),
                review_notes=r.review_notes if show and r is not None else None,
                awaiting_marking=bool(written and r is not None and r.marks_awarded is None),
            ))
        return AttemptReviewRead(
            attempt_id=attempt.id, exam_id=exam.id, exam_title=exam.title, show_result=show,
            total_marks=exam.total_marks,
            final_score=attempt.final_score if show else None,
            percentage=attempt.percentage if show else None,
            passed=attempt.passed if show else None,
            awaiting_marking=AttemptService.awaiting_marking(session, [attempt.id]).get(attempt.id, 0),
            items=out,
        )

    @staticmethod
    def get_exam_content(session: Session, attempt_id: UUID, student_id: UUID) -> AttemptExamRead:
        attempt = AttemptService.get_attempt(session, attempt_id, student_id)

        exam = session.exec(
            select(ExamModel).where(ExamModel.id == attempt.exam_id)
        ).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Exam not found.")

        sections_raw = ExamService.get_sections_with_items_internal(session, attempt.exam_id)

        all_item_ids = [
            item.item_id
            for section in sections_raw
            for item in section.items
        ]

        # item_bank_client = ItemBankClient()
        # items_by_id = await item_bank_client.get_items_for_display(all_item_ids)
        items = ItemService.get_by_ids_for_display(session, all_item_ids)
        # items_by_id = [ItemForDisplayRead.model_validate(i, from_attributes=True) for i in items]
        items_by_id = {
            item.id: ItemForDisplayRead.model_validate(item, from_attributes=True).model_dump()
            for item in items
        }

        sections = []
        for section in sections_raw:
            items = []
            for exam_item in section.items:
                item_id = exam_item.item_id
                content = items_by_id.get(item_id)
                if content is None:
                    continue
                items.append(AttemptItemRead(
                    id=item_id,
                    exam_item_id=exam_item.id,
                    question_text=content["question_text"],
                    item_type=content["item_type"],
                    options=content.get("options"),
                    marks=exam_item.marks,
                    word_limit=content.get("word_limit"),
                ))

            if exam.shuffle_options:
                items = AttemptService._shuffle_options(items, attempt_id)
            if exam.shuffle_questions:
                items = AttemptService._shuffle_items(items, attempt_id, section.id)

            sections.append(AttemptSectionRead(
                id=section.id,
                title=section.title,
                items=items,
            ))

        if exam.shuffle_questions:
            sections = AttemptService._shuffle_sections(sections, attempt_id)

        # The clock starts now, when the questions are first sent: time spent
        # loading the page on a slow connection doesn't count.
        now = datetime.now(timezone.utc)
        if attempt.status == AttemptStatus.STARTED and attempt.clock_started_at is None:
            attempt.clock_started_at = now
            session.add(attempt)
            session.commit()
            session.refresh(attempt)
        expires_at, remaining = AttemptService._timing(attempt, exam, now)

        return AttemptExamRead(
            attempt_id=attempt.id,
            exam_id=attempt.exam_id,
            duration_minutes=exam.duration_minutes,
            extra_time_seconds=exam.extra_time_seconds or 0,
            expires_at=expires_at,
            remaining_seconds=remaining,
            server_time=now,
            shuffle_questions=exam.shuffle_questions,
            shuffle_options=exam.shuffle_options,
            sections=sections,
        )

    

    # @staticmethod
    # def get_exam_content(session: Session, attempt_id: UUID, student_id: UUID) -> AttemptExamRead:
    #     attempt, exam = AttemptService.get_attempt_and_exam(session, attempt_id, student_id)

    #     sections_raw = ExamService.get_sections_with_items_internal(session, attempt.exam_id)
    #     all_item_ids: list[UUID] = [i.item_id for s in sections_raw for i in s.items]

    #     db_items = ItemService.get_by_ids_for_display(session, all_item_ids)
    #     items_by_id: dict[UUID, dict[str, Any]] = {
    #         it.id: ItemForDisplayRead.model_validate(it, from_attributes=True).model_dump()
    #         for it in db_items
    #     }

    #     sections: list[AttemptSectionRead] = []
    #     for section in sections_raw:
    #         section_items: list[AttemptItemRead] = []
    #         for exam_item in section.items:
    #             content = items_by_id.get(exam_item.item_id)
    #             if content is None:
    #                 continue
    #             section_items.append(AttemptItemRead(
    #                 id=exam_item.item_id,
    #                 exam_item_id=exam_item.id,
    #                 question_text=content["question_text"],
    #                 item_type=content["item_type"],
    #                 options=content.get("options"),
    #                 marks=exam_item.marks,
    #             ))
    #         if exam.shuffle_options:
    #             section_items = AttemptService._shuffle_options(section_items, attempt_id)
    #         sections.append(AttemptSectionRead(id=section.id, title=section.title, items=section_items))

    #     if exam.shuffle_questions:
    #         sections = AttemptService._shuffle_sections(sections, attempt_id)

    #     return AttemptExamRead(
    #         attempt_id=attempt.id, exam_id=attempt.exam_id,
    #         duration_minutes=exam.duration_minutes,
    #         shuffle_questions=exam.shuffle_questions,
    #         shuffle_options=exam.shuffle_options,
    #         sections=sections,
    #     )


    @staticmethod
    def _shuffle_sections(sections: list[AttemptSectionRead], attempt_id: UUID) -> list[AttemptSectionRead]:
        # Seed with attempt_id so the order is stable across repeated
        # GETs for the same attempt (autosave, page refresh, etc.)
        # but differs per student.
        rng = random.Random(str(attempt_id))
        shuffled = sections.copy()
        rng.shuffle(shuffled)
        return shuffled

    @staticmethod
    def _shuffle_items(items: list[AttemptItemRead], attempt_id: UUID, section_id: UUID) -> list[AttemptItemRead]:
        """Questions within a section, in a stable order per attempt (refreshes keep it)."""
        rng = random.Random(f"{attempt_id}:{section_id}")
        shuffled = items.copy()
        rng.shuffle(shuffled)
        return shuffled

    @staticmethod
    def _shuffle_options(items: list[AttemptItemRead], attempt_id: UUID) -> list[AttemptItemRead]:
        rng = random.Random(str(attempt_id))
        for item in items:
            if item.options and item.item_type not in FIXED_OPTION_TYPES:
                rng.shuffle(item.options)
        return items


    # @staticmethod
    # def get_cohort_attempts(
    #     session: Session,
    #     exam_id: UUID,
    #     cohort_id: UUID,
    #     current_user: CurrentUser,
    # ) -> list[CohortAttemptSummary]:
    #     attempts = session.exec(
    #         select(AttemptModel).where(
    #             AttemptModel.exam_id == exam_id,
    #             AttemptModel.org_id == current_user.org_id,
    #         )
    #     ).all()

    #     # Filter to the cohort via assignment lookup, since AttemptModel stores assignment_id
    #     assignment_ids = {a.assignment_id for a in attempts}
    #     cohort_assignments = session.exec(
    #         select(ExamAssignment.id).where(
    #             ExamAssignment.id.in_(assignment_ids),
    #             ExamAssignment.cohort_id == cohort_id,
    #         )
    #     ).all()
    #     cohort_assignment_ids = set(cohort_assignments)

    #     return [
    #         CohortAttemptSummary(
    #             id=a.id, student_id=a.student_id, status=a.status,
    #             submitted_at=a.submitted_at, final_score=a.final_score,
    #             percentage=a.percentage, passed=a.passed,
    #         )
    #         for a in attempts if a.assignment_id in cohort_assignment_ids
    #     ]
    @staticmethod
    def get_cohort_attempts(
        session: Session,
        exam_id: UUID,
        cohort_id: UUID | None,
        current_user: CurrentUser,
    ) -> list[CohortAttemptSummary]:
        """Attempts by students assigned through a cohort, or (cohort_id None) assigned one by one."""
        attempts: list[AttemptModel] = list(session.exec(
            AttemptService.cohort_attempts_stmt(session, exam_id, cohort_id, current_user)
        ).all())
        summaries: list[CohortAttemptSummary] = AttemptService.summaries(session, attempts)

        # Done with the DB: return the connection to the pool now, so it isn't
        # held while the router waits on the auth service.
        session.close()
        return summaries
    

    @staticmethod
    def cohort_attempts_stmt(
        session: Session, exam_id: UUID, cohort_id: UUID | None, current_user: CurrentUser,
    ) -> "SelectOfScalar[AttemptModel]":
        """Attempts by a cohort's students (cohort_id None: assigned one by one); checks access."""
        exam: ExamModel = ExamService._get_exam(session, exam_id, current_user.org_id)
        ExamService._assert_can_view(session, exam, current_user, P.RESULTS_VIEW)
        return (
            select(AttemptModel)
            .join(ExamAssignment, ExamAssignment.id == AttemptModel.assignment_id)
            .where(
                AttemptModel.exam_id == exam_id,
                AttemptModel.org_id == current_user.org_id,
                ExamAssignment.cohort_id == cohort_id if cohort_id is not None
                else col(ExamAssignment.cohort_id).is_(None),
            )
            .order_by(col(AttemptModel.submitted_at).desc().nulls_first(), col(AttemptModel.id))
        )

    @staticmethod
    def summaries(session: Session, attempts: list[AttemptModel]) -> list[CohortAttemptSummary]:
        waiting: dict[UUID, int] = AttemptService.awaiting_marking(session, [a.id for a in attempts])
        late: dict[UUID, int] = AttemptService.late_pending_counts(session, [a.id for a in attempts])
        return [
            CohortAttemptSummary(
                id=a.id, student_id=a.student_id, status=a.status, attempt_number=a.attempt_number,
                submitted_at=a.submitted_at, final_score=a.final_score,
                percentage=a.percentage, passed=a.passed,
                awaiting_marking=waiting.get(a.id, 0),
                late_answers=late.get(a.id, 0),
            )
            for a in attempts
        ]

    @staticmethod
    def attempts_page(
        session: Session, exam_id: UUID, cohort_id: UUID | None, current_user: CurrentUser,
        page: int, page_size: int,
    ) -> AttemptPage:
        """One page of attempts, with stats over all of them (fully marked ones only)."""
        stmt = AttemptService.cohort_attempts_stmt(session, exam_id, cohort_id, current_user)
        rows, total = page_of(session, stmt, page, page_size)
        ids_subq = stmt.order_by(None).with_only_columns(col(AttemptModel.id)).subquery()
        pending_ids = select(ResponseModel.attempt_id).where(col(ResponseModel.marks_awarded).is_(None))
        scored = session.exec(
            select(
                func.count(),
                func.avg(AttemptModel.percentage),
                func.count().filter(col(AttemptModel.passed).is_(True)),
            ).where(
                col(AttemptModel.id).in_(select(ids_subq)),
                col(AttemptModel.percentage).is_not(None),
                col(AttemptModel.id).not_in(pending_ids),
            )
        ).one()
        n_scored: int = int(scored[0] or 0)
        return AttemptPage(
            items=AttemptService.summaries(session, rows), total=total, page=page, page_size=page_size,
            total_pages=total_pages(total, page_size),
            scored=n_scored,
            mean_percentage=round(float(scored[1]), 2) if n_scored else None,
            pass_rate=round(int(scored[2] or 0) / n_scored * 100, 2) if n_scored else None,
        )

    @staticmethod
    def _get_items_for_scoring(session: Session, item_ids: list[UUID]) -> dict[UUID, dict[str, Any]]: # list[ItemForScoringRead]:
        items = ItemService.get_by_ids_for_scoring(session, item_ids)
        # return [ItemForScoringRead.model_validate(i, from_attributes=True) for i in items]
        return {
            item.id: ItemForScoringRead.model_validate(item, from_attributes=True).model_dump()
            for item in items
        }



    @staticmethod
    def get_attempt_detail_for_staff(
        session: Session,
        attempt_id: UUID,
        current_user: CurrentUser,
    ) -> AttemptDetailRead:
        attempt = session.exec(
            select(AttemptModel).where(
                AttemptModel.id == attempt_id,
                AttemptModel.org_id == current_user.org_id,
            )
        ).first()
        if not attempt:
            raise HTTPException(status_code=404, detail="Attempt not found.")
        ExamService._assert_can_view(
            session, ExamService._get_exam(session, attempt.exam_id, current_user.org_id),
            current_user, P.RESULTS_VIEW,
        )

        responses = session.exec(
            select(ResponseModel).where(ResponseModel.attempt_id == attempt_id)
        ).all()

        items: dict[UUID, ItemModel] = {
            i.id: i for i in session.exec(
                select(ItemModel).where(col(ItemModel.id).in_({r.item_id for r in responses}))
            ).all()
        } if responses else {}
        exam_items = AttemptService._exam_items(session, attempt.exam_id)
        order: dict[UUID, int] = {ei.item_id: ei.order for ei in exam_items.values()}

        response_details = []
        for r in sorted(responses, key=lambda r: order.get(r.item_id, 10**6)):
            item = items.get(r.item_id)
            exam_item = exam_items.get(r.item_id)
            response_details.append(AttemptResponseDetail(
                item_id=r.item_id,
                response_id=r.id,
                exam_item_id=exam_item.id if exam_item else r.exam_item_id,
                item_type=getattr(item.item_type, "value", None) if item else None,
                question_text=item.question_text if item else "[question no longer available]",
                answer=r.answer,
                correct_answers=item.correct_answers if item else None,
                is_correct=r.is_correct,
                marks_awarded=r.marks_awarded,
                max_marks=exam_item.marks if exam_item else (item.marks if item else None),
                marking_guide=item.marking_guide if item else None,
                review_notes=r.review_notes,
                reviewed_at=r.reviewed_at,
            ))

        return AttemptDetailRead(
            id=attempt.id, student_id=attempt.student_id, exam_id=attempt.exam_id,
            status=attempt.status, started_at=attempt.started_at, submitted_at=attempt.submitted_at,
            raw_score=attempt.raw_score, final_score=attempt.final_score,
            percentage=attempt.percentage, passed=attempt.passed,
            awaiting_marking=sum(1 for r in responses if r.marks_awarded is None),
            responses=response_details,
            late_answers=AttemptService.late_answers_for(session, attempt, items, list(responses)),
        )



    @staticmethod
    def reset_attempt(
        session: Session,
        attempt_id: UUID,
        payload: ResetAttemptRequest,
        current_user: CurrentUser,
    ) -> AttemptModel:
        # Lock the row so a violation or autosave arriving mid-reset can't
        # interleave with the clean-up below.
        attempt: AttemptModel | None = session.exec(
            select(AttemptModel)
            .where(
                AttemptModel.id == attempt_id,
                AttemptModel.org_id == current_user.org_id,
            )
            .with_for_update()
        ).first()
        if not attempt:
            raise HTTPException(status_code=404, detail="Attempt not found.")

        # 1. Wipe responses (and any late answers) so the student gets a clean slate.
        session.exec(
            delete(ResponseModel).where(ResponseModel.attempt_id == attempt_id)
        )
        session.exec(
            delete(LateResponseModel).where(LateResponseModel.attempt_id == attempt_id)
        )

        # 2. Soft-clear violations: keep the rows as evidence, but stop them
        #    counting toward max_violations. Already-cleared rows keep their
        #    original cleared_at.
        active_violation_filter = (
            ExamViolation.attempt_id == attempt_id,
            col(ExamViolation.cleared_at).is_(None),
        )
        violations_cleared: int = session.exec(
            select(func.count()).select_from(ExamViolation).where(*active_violation_filter)
        ).one()

        if violations_cleared:
            session.exec(
                update(ExamViolation)
                .where(*active_violation_filter)
                .values(cleared_at=datetime.now(timezone.utc))
            )

        # 3. Reset the attempt itself. start() reactivates this same row.
        attempt.status = AttemptStatus.RESET
        attempt.submitted_at = None
        attempt.raw_score = None
        attempt.final_score = None
        attempt.percentage = None
        attempt.passed = None
        attempt.scored_at = None
        attempt.scored_by = None
        session.add(attempt)
        AttemptService._set_assignment_status(session, attempt, AssignmentStatus.ASSIGNED)

        # 4. Optionally reschedule the assignment.
        if payload.new_scheduled_at:
            assignment: ExamAssignment | None = session.exec(
                select(ExamAssignment).where(ExamAssignment.id == attempt.assignment_id)
            ).first()
            if assignment:
                exam: ExamModel | None = session.exec(
                    select(ExamModel).where(ExamModel.id == attempt.exam_id)
                ).first()
                assignment.scheduled_at = payload.new_scheduled_at
                assignment.deadline_override = (
                    payload.new_scheduled_at + AttemptService._allowed(exam)
                    if exam else None
                )
                session.add(assignment)

        # 5. Audit trail, including how many violations were cleared.
        _log(
            session,
            attempt.exam_id,
            current_user.org_id,
            current_user.id,
            "attempt_reset",
            {
                "attempt_id": str(attempt_id),
                "reason": payload.reason,
                "violations_cleared": violations_cleared,
            },
        )

        # Single commit: responses, violations, attempt, assignment and log
        # all succeed or all roll back together.
        session.commit()
        session.refresh(attempt)
        return attempt


    # @staticmethod
    # def reset_attempt(
    #     session: Session,
    #     attempt_id: UUID,
    #     payload: ResetAttemptRequest,
    #     current_user: CurrentUser,
    # ) -> AttemptModel:
    #     attempt = session.exec(
    #         select(AttemptModel).where(
    #             AttemptModel.id == attempt_id,
    #             AttemptModel.org_id == current_user.org_id,
    #         )
    #     ).first()
    #     if not attempt:
    #         raise HTTPException(status_code=404, detail="Attempt not found.")

    #     # Wipe responses so the student gets a clean slate
    #     session.exec(
    #         delete(ResponseModel).where(ResponseModel.attempt_id == attempt_id)
    #     )

    #     attempt.status = AttemptStatus.RESET
    #     attempt.submitted_at = None
    #     attempt.raw_score = None
    #     attempt.final_score = None
    #     attempt.percentage = None
    #     attempt.passed = None
    #     attempt.scored_at = None
    #     attempt.scored_by = None
    #     session.add(attempt)

    #     if payload.new_scheduled_at:
    #         assignment = session.exec(
    #             select(ExamAssignment).where(ExamAssignment.id == attempt.assignment_id)
    #         ).first()
    #         if assignment:
    #             exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
    #             assignment.scheduled_at = payload.new_scheduled_at
    #             assignment.deadline_override = (
    #                 payload.new_scheduled_at + timedelta(minutes=exam.duration_minutes)
    #                 if exam else None
    #             )
    #             session.add(assignment)

    #     _log(session, attempt.exam_id, current_user.org_id, current_user.id,
    #         "attempt_reset", {"attempt_id": str(attempt_id), "reason": payload.reason})

    #     session.commit()
    #     session.refresh(attempt)
    #     return attempt





    @staticmethod
    def awaiting_marking(session: Session, attempt_ids: list[UUID]) -> dict[UUID, int]:
        """Per attempt: answers still waiting for a person to mark them."""
        if not attempt_ids:
            return {}
        rows = session.exec(
            select(ResponseModel.attempt_id, func.count())
            .where(col(ResponseModel.attempt_id).in_(attempt_ids), col(ResponseModel.marks_awarded).is_(None))
            .group_by(ResponseModel.attempt_id)
        ).all()
        return {attempt_id: int(n) for attempt_id, n in rows}

    @staticmethod
    def attempt_of(session: Session, response: ResponseModel) -> AttemptModel | None:
        return session.get(AttemptModel, response.attempt_id)

    @staticmethod
    def _exam_items(session: Session, exam_id: UUID) -> dict[UUID, ExamItem]:
        """The exam's questions keyed by ITEM id (one entry per question per exam)."""
        return {
            ei.item_id: ei for ei in session.exec(select(ExamItem).where(ExamItem.exam_id == exam_id)).all()
        }

    @staticmethod
    def recompute_score(
        session: Session,
        attempt: AttemptModel,
        exam: ExamModel,
        responses: list[ResponseModel] | None = None,
        pass_mark: float | None = None,
    ) -> int:
        """
        Totals, percentage and pass/fail from the answers' marks. Out of the
        whole exam (questions left unanswered count as 0), not just the ones
        answered. Returns how many answers still wait for a person to mark;
        pass/fail is only decided once that's 0. Doesn't commit.
        """
        if responses is None:
            responses = list(session.exec(
                select(ResponseModel).where(ResponseModel.attempt_id == attempt.id)
            ).all())
        total_possible: float = sum(ei.marks for ei in AttemptService._exam_items(session, exam.id).values())
        if total_possible <= 0:
            total_possible = exam.total_marks or 0.0

        on_exam = AttemptService._exam_items(session, exam.id)
        responses = [r for r in responses if r.item_id in on_exam]
        awarded: list[float] = [r.marks_awarded for r in responses if r.marks_awarded is not None]
        pending: int = sum(1 for r in responses if r.marks_awarded is None)

        attempt.raw_score = sum(m for m in awarded if m > 0)
        attempt.final_score = max(0.0, sum(awarded))
        attempt.percentage = round(attempt.final_score / total_possible * 100, 2) if total_possible else 0.0
        attempt.scored_at = datetime.now(timezone.utc)

        threshold = pass_threshold(pass_mark if pass_mark is not None else exam.pass_mark)
        attempt.passed = None if pending or threshold is None else attempt.percentage >= threshold - 1e-9
        session.add(attempt)
        return pending

    @staticmethod
    def _finalize_submission(
        session: Session,
        attempt: AttemptModel,
        exam: ExamModel,
        pass_mark: float | None = None,
        auto: bool = False,
        final_status: AttemptStatus = AttemptStatus.SUBMITTED,   # NEW
    ) -> AttemptModel:
        """
        Scores and closes out an attempt. Shared by student-triggered submit,
        lazy time-expiry auto-submit, and violation-triggered termination —
        all three score identically; only the resulting status differs.
        """
        responses = session.exec(
            select(ResponseModel).where(ResponseModel.attempt_id == attempt.id)
        ).all()

        item_ids = [r.item_id for r in responses]
        item_bank = AttemptService._get_items_for_scoring(session, item_ids)
        exam_items = AttemptService._exam_items(session, exam.id)

        for resp in responses:
            item_data = item_bank.get(resp.item_id)
            exam_item = exam_items.get(resp.item_id)
            if not item_data or exam_item is None:
                # Not a question on this exam: never scored.
                resp.is_correct = None
                resp.marks_awarded = 0.0
                session.add(resp)
                continue
            # The exam's snapshot of the marks, so editing a question later
            # doesn't change an exam that has already run.
            marks: float = exam_item.marks if exam_item else item_data.get("marks", 1.0)
            neg: float = exam_item.negative_marks if exam_item else item_data.get("negative_marks", 0.0)
            score = score_answer(
                str(getattr(item_data.get("item_type"), "value", item_data.get("item_type"))),
                item_data.get("correct_answers"), resp.answer, marks, neg,
                auto_mark_short_answers=bool(exam.auto_mark_short_answers),
            )
            resp.is_correct = score.is_correct
            resp.marks_awarded = score.marks_awarded
            session.add(resp)

        attempt.status = final_status
        now = datetime.now(timezone.utc)
        # Closed because time ran out (possibly noticed later): record the deadline itself.
        deadline = AttemptService._deadline(attempt, exam)
        attempt.submitted_at = min(now, deadline) if auto and deadline is not None else now
        attempt.scored_by = (
            "violation_termination" if final_status == AttemptStatus.TERMINATED
            else "auto_expired" if auto
            else "auto"
        )
        AttemptService.recompute_score(session, attempt, exam, list(responses), pass_mark=pass_mark)

        session.add(attempt)
        AttemptService._set_assignment_status(session, attempt, AssignmentStatus.SUBMITTED)
        session.commit()
        session.refresh(attempt)
        return attempt



    @staticmethod
    def _enforce_time_limit(session: Session, attempt: AttemptModel) -> AttemptModel:
        if attempt.status != AttemptStatus.STARTED:
            return attempt

        exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
        if not exam:
            return attempt

        deadline = AttemptService._deadline(attempt, exam)
        if deadline is None or datetime.now(timezone.utc) < deadline:
            return attempt

        return AttemptService._finalize_submission(session, attempt, exam, auto=True)


    @staticmethod
    def _enforce_time_limit_new(
        session: Session, attempt: AttemptModel, exam: ExamModel
    ) -> AttemptModel:
        if attempt.status != AttemptStatus.STARTED:
            return attempt
        deadline = AttemptService._deadline(attempt, exam)
        if deadline is None or datetime.now(timezone.utc) < deadline:
            return attempt
        return AttemptService._finalize_submission(session, attempt, exam, auto=True)


    @staticmethod
    def _to_attempt_read(session: Session, attempt: AttemptModel, for_student: bool = False) -> AttemptRead:
        """
        for_student: the candidate's own view. If the exam doesn't show results,
        scores and pass/fail are left out.
        """
        exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
        show_result: bool = bool(exam.show_result_immediately) if exam else True
        hide: bool = for_student and not show_result
        now = datetime.now(timezone.utc)

        expires_at, remaining = AttemptService._timing(attempt, exam, now)

        return AttemptRead(
            id=attempt.id, exam_id=attempt.exam_id, assignment_id=attempt.assignment_id, student_id=attempt.student_id,
            status=attempt.status, attempt_number=attempt.attempt_number,
            started_at=attempt.started_at, submitted_at=attempt.submitted_at,
            raw_score=None if hide else attempt.raw_score,
            final_score=None if hide else attempt.final_score,
            percentage=None if hide else attempt.percentage,
            passed=None if hide else attempt.passed,
            expires_at=expires_at, remaining_seconds=remaining, server_time=now,
            show_result=show_result, allow_review=bool(exam.allow_review) if exam else False,
            clock_started=attempt.clock_started_at is not None,
            awaiting_marking=(
                0 if attempt.status == AttemptStatus.STARTED
                else AttemptService.awaiting_marking(session, [attempt.id]).get(attempt.id, 0)
            ),
        )

    @staticmethod
    def _to_attempt_read_new(attempt: AttemptModel, exam: ExamModel | None) -> AttemptRead:
        now = datetime.now(timezone.utc)
        expires_at, remaining = AttemptService._timing(attempt, exam, now)
        return AttemptRead(
            id=attempt.id, exam_id=attempt.exam_id, assignment_id=attempt.assignment_id,
            student_id=attempt.student_id, status=attempt.status,
            attempt_number=attempt.attempt_number, started_at=attempt.started_at,
            submitted_at=attempt.submitted_at, raw_score=attempt.raw_score,
            final_score=attempt.final_score, percentage=attempt.percentage,
            passed=attempt.passed, expires_at=expires_at,
            remaining_seconds=remaining, server_time=now,
            clock_started=attempt.clock_started_at is not None,
        )


    @staticmethod
    def sweep_expired_attempts(session: Session) -> int:
        now = datetime.now(timezone.utc)
        candidates = session.exec(
            select(AttemptModel).where(AttemptModel.status == AttemptStatus.STARTED)
        ).all()
        expired = 0
        for attempt in candidates:
            exam = session.exec(select(ExamModel).where(ExamModel.id == attempt.exam_id)).first()
            if not exam:
                continue
            deadline = AttemptService._deadline(attempt, exam)
            if deadline is not None and now >= deadline:
                asyncio.run(AttemptService._finalize_submission(session, attempt, exam, auto=True))
                expired += 1
        return expired








    @staticmethod
    def load_exam_export_data(
        session: Session,
        exam_id: UUID,
        org_id: UUID,
        current_user: CurrentUser | None = None,
    ) -> ExamExportData:
        """Load the exam, attempts, and responses using SQLModel."""

        exam = session.exec(
            select(ExamModel).where(
                ExamModel.id == exam_id,
                ExamModel.org_id == org_id,
            )
        ).first()

        if exam is None:
            raise HTTPException(
                status_code=404,
                detail="Exam not found.",
            )
        if current_user is not None:
            ExamService._assert_can_view(session, exam, current_user, P.RESULTS_EXPORT)

        sections = ExamService.get_sections_with_items_internal(
            session,
            exam_id,
        )

        exam_items = [
            item
            for section in sections
            for item in section.items
        ]

        attempts = list(
            session.exec(
                select(AttemptModel).where(
                    AttemptModel.exam_id == exam_id,
                    AttemptModel.org_id == org_id,
                )
            ).all()
        )

        if not attempts:
            raise HTTPException(
                status_code=404,
                detail="No attempts found for this exam.",
            )

        attempt_ids = [attempt.id for attempt in attempts]

        responses = list(
            session.exec(
                select(ResponseModel).where(
                    ResponseModel.attempt_id.in_(attempt_ids),
                    ResponseModel.org_id == org_id,
                )
            ).all()
        )

        return ExamExportData(
            exam_items=exam_items,
            attempts=attempts,
            responses=responses,
        )


def answer_is_present(answer: Any) -> bool:
    """Check for an answer without rejecting valid 0 or False values."""

    if answer is None:
        return False

    if isinstance(answer, str):
        return bool(answer.strip())

    if isinstance(answer, (list, dict, tuple, set)):
        return len(answer) > 0

    return True


def build_exam_attempts_workbook(
    export_data: ExamExportData,
    users_by_id: dict,
) -> BytesIO:
    """Create the Excel workbook from loaded exam and student data."""

    exam_items = export_data.exam_items
    attempts = export_data.attempts
    responses = export_data.responses

    # Group responses by attempt.
    responses_by_attempt: dict[UUID, list[ResponseModel]] = {
        attempt.id: [] for attempt in attempts
    }

    for response in responses:
        responses_by_attempt.setdefault(
            response.attempt_id, []
        ).append(response)

    # Sort by score, highest first. Missing scores rank below scored attempts.
    ranked_attempts = sorted(
        attempts,
        key=lambda attempt: (
            attempt.final_score is not None,
            (
                float(attempt.final_score)
                if attempt.final_score is not None
                else 0.0
            ),
        ),
        reverse=True,
    )

    # Competition ranking: 1, 1, 3 for tied first-place scores.
    ranks: dict[UUID, int] = {}
    previous_score = object()
    current_rank = 0

    for position, attempt in enumerate(ranked_attempts, start=1):
        score = attempt.final_score

        if score != previous_score:
            current_rank = position
            previous_score = score

        ranks[attempt.id] = current_rank

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Exam Results"

    headers = [
        "Name",
        "Enrollment No.",
        "Email address",
        "State",
        "Started on",
        "Completed",
        "Time taken",
        "Total Questions",
        "Questions Attempted",
        "Questions Correct",
        "Questions Incorrect",
        "Percentage",
        "Rank",
        "Grade/20.00",
    ]

    for index, exam_item in enumerate(exam_items, start=1):
        marks = getattr(exam_item, "marks", None)
        marks = float(marks) if marks is not None else 1.0
        headers.append(f"Q. {index} /{marks:.2f}")

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

    for attempt in ranked_attempts:
        # AuthClient returns UserSummary objects, not dictionaries.
        student = users_by_id.get(str(attempt.student_id))

        if student:
            student_name = " ".join(
                value
                for value in (
                    student.firstname,
                    student.lastname,
                )
                if value
            ).strip()
            email_address = student.email or ""
        else:
            student_name = ""
            email_address = ""

        if not student_name:
            student_name = str(attempt.student_id)

        attempt_responses = responses_by_attempt.get(
            attempt.id, []
        )

        response_by_exam_item = {
            response.exam_item_id: response
            for response in attempt_responses
            if response.exam_item_id is not None
        }

        # Legacy fallback; safe only if each underlying item appears once
        # in the exam when exam_item_id is absent.
        response_by_item = {
            response.item_id: response
            for response in attempt_responses
        }

        question_results = []
        questions_attempted = 0
        questions_correct = 0
        questions_incorrect = 0

        for exam_item in exam_items:
            response = response_by_exam_item.get(exam_item.id)

            if response is None:
                response = response_by_item.get(exam_item.item_id)

            if response is None or not answer_is_present(
                response.answer
            ):
                question_results.append("-")
                continue

            questions_attempted += 1

            if response.is_correct is True:
                questions_correct += 1
                question_results.append(1)
            elif response.is_correct is False:
                questions_incorrect += 1
                question_results.append(0)
            else:
                # Answer exists but correctness is not yet determined.
                question_results.append("-")

        # Time taken counts from when the clock started (questions sent).
        started_at = attempt.clock_started_at or attempt.started_at
        submitted_at = attempt.submitted_at

        started_on = (
            started_at.strftime("%Y-%m-%d %H:%M:%S")
            if started_at
            else ""
        )

        completed = (
            submitted_at.strftime("%Y-%m-%d %H:%M:%S")
            if submitted_at
            else ""
        )

        time_taken = ""

        if started_at:
            end_time = submitted_at or datetime.now(
                started_at.tzinfo
            )

            elapsed_seconds = max(
                0,
                int((end_time - started_at).total_seconds()),
            )
            minutes, seconds = divmod(elapsed_seconds, 60)
            time_taken = f"{minutes}m {seconds}s"

        percentage = (
            float(attempt.percentage)
            if attempt.percentage is not None
            else 0.0
        )
        grade = round(percentage / 5, 2)

        worksheet.append(
            [
                student_name,
                "",  # Populate from the actual enrollment-number source.
                email_address,
                "",  # Populate if student state is available.
                started_on,
                completed,
                time_taken,
                len(exam_items),
                questions_attempted,
                questions_correct,
                questions_incorrect,
                percentage,
                ranks.get(attempt.id, ""),
                grade,
                *question_results,
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    fixed_widths = {
        "A": 28,
        "B": 24,
        "C": 32,
        "D": 16,
        "E": 22,
        "F": 22,
        "G": 16,
        "H": 16,
        "I": 22,
        "J": 20,
        "K": 22,
        "L": 14,
        "M": 10,
        "N": 14,
    }

    for column, width in fixed_widths.items():
        worksheet.column_dimensions[column].width = width

    for column_index in range(15, worksheet.max_column + 1):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = 14

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output
