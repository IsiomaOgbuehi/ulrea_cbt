import jwt as pyjwt

import pytest
from io import BytesIO
from openpyxl import Workbook
from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import Workbook, load_workbook 
from typing import Any
from uuid import UUID, uuid4
from datetime import datetime
from unittest.mock import AsyncMock, patch
from fastapi import HTTPException

from sqlmodel import select, Session
from ..conftest import engine

from auth.database.schema.cohort.cohort_db import CohortModel, CohortMember
from auth.database.schema.user.user_db import UserModel
from auth.database.schema.membership.membership_db import OrgMembership

from auth.database.schema.user.enums import MembershipStatus
from auth.services.user.user_management_service import UserManagementService
from auth.utility.password.password_harsher import PasswordHasher

from tests.conftest import (
    do_signup, do_request_otp, do_verify_otp,
    do_full_signup, SIGNUP_PAYLOAD, USER_EMAIL, make_staff_activation_token
)


ADMIN_PAYLOAD = {
    "firstname": "Alice",
    "lastname": "Admin",
    "email": "alice@cbtech.com",
    "phone": "+1000000001",
    "role": "admin",
}

TEACHER_PAYLOAD = {
    "firstname": "Bob",
    "lastname": "Teacher",
    "email": "bob@cbtech.com",
    "phone": "+1000000002",
    "role": "teacher",
}

STUDENT_PAYLOAD = {
    "firstname": "Charlie",
    "lastname": "Student",
    "phone": "+1000000003",
    "institution_id": "STU/2024/001",   # required reg number
}


def create_student_excel(rows):
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append([
        "firstname",
        "lastname",
        "othername",
        "email",
        "phone",
        "institution_id",
        "access_code",
    ])

    for row in rows:
        worksheet.append([
            row.get("firstname"),
            row.get("lastname"),
            row.get("othername"),
            row.get("email"),
            row.get("phone"),
            row.get("institution_id"),
            row.get("access_code"),
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def get_super_admin_token(client) -> str:
    verify_data = do_full_signup(client)
    return verify_data['token']['access_token']


def _activate_staff(
    client,
    user_id: str,
    email: str,
    password: str = "newSecurePass123!",
) -> dict:
    activation_token = make_staff_activation_token(user_id, email)

    resp = client.post(
        "/api/v1/users/staff/activate",
        json={
            "token": activation_token,
            "password": password,
            "confirm_password": password,
        },
    )

    assert resp.status_code == 200, resp.json()
    return resp.json()


# def _activate_staff(client, user_id: str, password: str = "newSecurePass123!") -> dict:
#     """Helper: generate activation token and complete staff activation flow."""
#     activation_token = make_staff_activation_token(user_id)
#     resp = client.post(
#         "/api/v1/users/staff/activate",
#         json={
#             "token": activation_token,
#             "password": password,
#             "confirm_password": password,
#         },
#     )
#     assert resp.status_code == 200, resp.json()
#     return resp.json()


# def _activate_staff(client, user_id: str, password: str = "newSecurePass123!") -> dict:
#     """Helper: generate activation token and complete staff activation flow."""
#     from auth.utility.jwt.token_activation import create_staff_activation_token
#     activation_token = create_staff_activation_token(user_id)
#     resp = client.post(
#         "/api/v1/users/staff/activate",
#         json={
#             "token": activation_token,
#             "password": password,
#             "confirm_password": password,
#         },
#     )
#     assert resp.status_code == 200, resp.json()
#     return resp.json()



def _create_pending_teacher(client, token: str) -> dict:
    resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 200, resp.json()
    return resp.json()


# ============================================================
# STAFF MANAGEMENT
# ============================================================

def test_super_admin_can_create_admin(client):
    token = get_super_admin_token(client)

    response = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200, response.json()
    data = response.json()
    assert data['role'] == 'admin'
    assert data['is_first_login'] is True
    assert 'temporary_password' in data


def test_super_admin_can_create_student(client):
    token = get_super_admin_token(client)

    response = client.post(
        "/api/v1/users/students/create",
        json=STUDENT_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200, response.json()
    data = response.json()
    assert data['role'] == 'student'
    assert len(data['access_code']) == 11
    assert data['is_first_login'] is True


def test_admin_can_create_teacher(client):
    """Smoke test: super admin creates an admin, admin activates, admin creates teacher."""
    super_token = get_super_admin_token(client)

    # Create admin
    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    assert create_resp.status_code == 200, create_resp.json()
    admin_id = create_resp.json()["id"]

    # Admin activates account
    activate_data = _activate_staff(client, admin_id, ADMIN_PAYLOAD["email"])
    admin_token = activate_data["access_token"]

    # Admin creates a teacher
    response = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert response.status_code == 200, response.json()
    assert response.json()["role"] == "teacher"


def test_staff_first_login_setup(client):
    """Staff creation → activation token flow (temp-password login is deprecated)."""
    token = get_super_admin_token(client)

    create_response = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_response.status_code == 200, create_response.json()
    created_user = create_response.json()

    # Activate via token — temp-password login is no longer supported
    activate_data = _activate_staff(client, created_user["id"], created_user["email"])
    assert "access_token" in activate_data

    # Login with newly set password works
    final_login = client.post(
        "/api/v1/auth/login",
        data={"username": TEACHER_PAYLOAD["email"], "password": "newSecurePass123!"},
    )
    assert final_login.status_code == 200, final_login.json()


def test_staff_account_activation(client):

    token = get_super_admin_token(client)

    create_response = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_response.status_code == 200, create_response.json()
    created_user = create_response.json()
    assert created_user["is_first_login"] is True

    activation_token = make_staff_activation_token(created_user["id"], TEACHER_PAYLOAD["email"])

    activate_response = client.post(
        "/api/v1/users/staff/activate",
        json={
            "token": activation_token,
            "password": "newSecurePass123!",
            "confirm_password": "newSecurePass123!",
        },
    )
    assert activate_response.status_code == 200, activate_response.json()
    activate_data = activate_response.json()
    assert activate_data["detail"] == "Account activated successfully."
    assert "access_token" in activate_data
    assert "refresh_token" in activate_data

    # Login with new password
    login_response = client.post(
        "/api/v1/auth/login",
        data={"username": TEACHER_PAYLOAD["email"], "password": "newSecurePass123!"},
    )
    assert login_response.status_code == 200, login_response.json()
    assert "access_token" in login_response.json()


def test_teacher_cannot_create_users(client):
    super_token = get_super_admin_token(client)

    # Create and activate teacher
    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    assert create_resp.status_code == 200, create_resp.json()
    teacher_id = create_resp.json()["id"]

    activate_data = _activate_staff(client, teacher_id, TEACHER_PAYLOAD["email"])
    teacher_token = activate_data["access_token"]

    # Teacher tries to create another user — role check fires before anything else
    response = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {teacher_token}"},
    )
    assert response.status_code == 403, response.json()


# ============================================================
# STUDENT MANAGEMENT
# ============================================================

def test_student_first_login_setup_and_login(client):
    token = get_super_admin_token(client)

    create_response = client.post(
        "/api/v1/users/students/create",
        json=STUDENT_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_response.status_code == 200, create_response.json()
    access_code = create_response.json()['access_code']

    # Student completes first login setup
    setup_response = client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": access_code,
            "favorite_question": "What is your pet's name?",
            "favorite_answer": "Fluffy",
        }
    )
    assert setup_response.status_code == 200, setup_response.json()
    assert 'access_token' in setup_response.json()

    # Student logs in with access code + favorite answer
    login_response = client.post(
        "/api/v1/users/student/login",
        json={
            "access_code": access_code,
            "favorite_answer": "Fluffy",
        }
    )
    assert login_response.status_code == 200, login_response.json()
    assert 'access_token' in login_response.json()


def test_student_cannot_use_staff_login(client):
    token = get_super_admin_token(client)

    client.post(
        "/api/v1/users/students/create",
        json=STUDENT_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )

    # Students have no email/password — staff login should reject them
    response = client.post(
        "/api/v1/auth/login",
        data={"username": "notanemail@x.com", "password": "anything"}
    )
    assert response.status_code == 401


def test_student_can_fetch_security_question(client):
    token = get_super_admin_token(client)

    create_response = client.post(
        "/api/v1/users/students/create",
        json=STUDENT_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_response.status_code == 200, create_response.json()
    access_code = create_response.json()["access_code"]

    # Complete first-time setup
    setup_response = client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": access_code,
            "favorite_question": "What is your pet's name?",
            "favorite_answer": "Fluffy",
        },
    )
    assert setup_response.status_code == 200, setup_response.json()

    # Fetch security question
    response = client.post(
        "/api/v1/users/student/login/question",
        json={"access_code": access_code},
    )
    assert response.status_code == 200, response.json()
    assert response.json()["favorite_question"] == "What is your pet's name?"


def test_student_question_invalid_access_code(client):
    response = client.post(
        "/api/v1/users/student/login/question",
        json={"access_code": "INVALID123"},
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "Invalid access code."


def test_student_question_requires_first_time_setup(client):
    token = get_super_admin_token(client)

    create_response = client.post(
        "/api/v1/users/students/create",
        json=STUDENT_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_response.status_code == 200, create_response.json()
    access_code = create_response.json()["access_code"]

    # No setup done yet — question endpoint must reject
    response = client.post(
        "/api/v1/users/student/login/question",
        json={"access_code": access_code},
    )
    assert response.status_code == 403
    assert response.json()["detail"] == "Please complete first-time setup first."


# ============================================================
# DEBUG (keep until redis path confirmed stable, then delete)
# ============================================================

def test_debug_redis_path(client):
    from auth.dependencies import auth_dependencies
    import auth.api.v1.routes.users as users_module

    token = get_super_admin_token(client)

    print("\nauth_dependencies redis_client:", auth_dependencies.redis_client)
    print("users module redis_client:", getattr(users_module, 'redis_client', 'NOT FOUND'))

    response = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    print("Response:", response.json())
    print("Status:", response.status_code)












# --------

def _decode_user_id(token: str) -> str:
    """Test-only helper — pulls `sub` out of a token without verifying signature."""
    payload = pyjwt.decode(token, options={"verify_signature": False})
    return payload["sub"]


# ============================================================
# ADMIN UPDATE USER
# ============================================================

def test_admin_can_update_staff_email(client):
    token = get_super_admin_token(client)

    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert create_resp.status_code == 200, create_resp.json()
    teacher_id = create_resp.json()["id"]

    update_resp = client.patch(
        f"/api/v1/users/staff/{teacher_id}",
        json={"email": "corrected@cbtech.com"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert update_resp.status_code == 200, update_resp.json()
    assert update_resp.json()["email"] == "corrected@cbtech.com"


def test_cannot_update_to_duplicate_email(client):
    token = get_super_admin_token(client)

    client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    teacher_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    teacher_id = teacher_resp.json()["id"]

    response = client.patch(
        f"/api/v1/users/staff/{teacher_id}",
        json={"email": ADMIN_PAYLOAD["email"]},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 409, response.json()


def test_teacher_cannot_update_other_users(client):
    token = get_super_admin_token(client)

    teacher_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    teacher_id = teacher_resp.json()["id"]
    teacher_token = _activate_staff(client, teacher_id, TEACHER_PAYLOAD["email"])["access_token"]

    another_resp = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    another_id = another_resp.json()["id"]

    response = client.patch(
        f"/api/v1/users/staff/{another_id}",
        json={"firstname": "Hacked"},
        headers={"Authorization": f"Bearer {teacher_token}"},
    )
    assert response.status_code == 403, response.json()


# ============================================================
# ADMIN DELETE USER
# ============================================================

def test_admin_can_delete_never_activated_staff(client):
    token = get_super_admin_token(client)

    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    teacher_id = create_resp.json()["id"]

    delete_resp = client.delete(
        f"/api/v1/users/{teacher_id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert delete_resp.status_code == 200, delete_resp.json()
    assert delete_resp.json()["action"] == "deleted"

    # confirm gone
    followup = client.patch(
        f"/api/v1/users/staff/{teacher_id}",
        json={"firstname": "Ghost"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert followup.status_code == 404


def test_cannot_delete_active_user_without_force(client):
    token = get_super_admin_token(client)

    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {token}"},
    )
    teacher_id = create_resp.json()["id"]
    _activate_staff(client, teacher_id, TEACHER_PAYLOAD["email"])  # now activated

    response = client.delete(
        f"/api/v1/users/{teacher_id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 400
    assert "force" in response.json()["detail"].lower()


def test_admin_cannot_force_delete_active_user(client):
    """force=true requires super_admin — a plain admin gets 403 even with the flag."""
    super_token = get_super_admin_token(client)

    admin_create = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    admin_id = admin_create.json()["id"]
    admin_token = _activate_staff(client, admin_id, ADMIN_PAYLOAD["email"])["access_token"]

    teacher_create = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    teacher_id = teacher_create.json()["id"]
    _activate_staff(client, teacher_id, TEACHER_PAYLOAD["email"])  # now activated

    response = client.delete(
        f"/api/v1/users/{teacher_id}?force=true",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert response.status_code == 403


def test_super_admin_can_force_delete_active_user(client):
    super_token = get_super_admin_token(client)

    create_resp = client.post(
        "/api/v1/users/staff/create",
        json=TEACHER_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    teacher_id = create_resp.json()["id"]
    _activate_staff(client, teacher_id, TEACHER_PAYLOAD["email"])  # now activated

    response = client.delete(
        f"/api/v1/users/{teacher_id}?force=true",
        headers={"Authorization": f"Bearer {super_token}"},
    )
    assert response.status_code == 200
    assert response.json()["action"] == "deleted"


def test_cannot_delete_super_admin(client):
    token = get_super_admin_token(client)
    super_admin_id = _decode_user_id(token)

    response = client.delete(
        f"/api/v1/users/{super_admin_id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.json()['detail'] == 'You cannot delete your own account.'
    # assert response.status_code == 403


def test_cannot_delete_self(client):
    super_token = get_super_admin_token(client)

    admin_create = client.post(
        "/api/v1/users/staff/create",
        json=ADMIN_PAYLOAD,
        headers={"Authorization": f"Bearer {super_token}"},
    )
    admin_id = admin_create.json()["id"]
    admin_token = _activate_staff(client, admin_id, ADMIN_PAYLOAD["email"])["access_token"]

    response = client.delete(
        f"/api/v1/users/{admin_id}",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert response.status_code == 400
    # assert "yourself" in response.json()["detail"].lower()


# Bulk create without cohort
def test_bulk_create_students_without_cohort(client):
    token = get_super_admin_token(client)

    excel_file = create_student_excel([
        {
            "firstname": "John",
            "lastname": "Doe",
            "email": "john.bulk@example.com",
            "institution_id": "848348349",
        },
        {
            "firstname": "Jane",
            "lastname": "Doe",
            "email": "jane.bulk@example.com",
            "institution_id": "848390696",
        },
    ])

    response = client.post(
        "/api/v1/users/students/create/bulk",
        files={
            "file": (
                "students.xlsx",
                excel_file.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={
            "Authorization": f"Bearer {token}",
        },
    )

    assert response.status_code == 200, response.json()

    data = response.json()

    print("\nBULK RESPONSE:")
    print(data)

    assert data["total_rows"] == 2
    assert data["successful_rows"] == 2
    assert data["failed_rows"] == 0
    assert len(data["students"]) == 2

# Bulk create with Cohort
def test_bulk_create_students_assigns_to_cohort(client):
    token = get_super_admin_token(client)

    # Get the authenticated user's organization
    with Session(engine) as session:
        user = session.exec(
            select(UserModel).where(
                UserModel.email == USER_EMAIL
            )
        ).first()

        assert user is not None

        membership = session.exec(
            select(OrgMembership).where(
                OrgMembership.user_id == user.id
            )
        ).first()

        assert membership is not None

        user_id = user.id
        org_id = membership.org_id

        cohort = CohortModel(
            org_id=org_id,
            name="JSS3A",
            description="Test cohort",
            created_by=user_id,
        )

        session.add(cohort)
        session.commit()
        session.refresh(cohort)

        cohort_id = cohort.id

    excel_file = create_student_excel([
        {
            "firstname": "John",
            "lastname": "Doe",
            "email": "john.cohort@example.com",
            "institution_id": "848398349",
        },
        {
            "firstname": "Jane",
            "lastname": "Doe",
            "email": "jane.cohort@example.com",
            "institution_id": "4309434",
        },
    ])

    response = client.post(
        f"/api/v1/users/students/create/bulk?cohort_id={cohort_id}",
        files={
            "file": (
                "students.xlsx",
                excel_file.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={
            "Authorization": f"Bearer {token}",
        },
    )

    assert response.status_code == 200, response.json()

    data = response.json()

    assert data["total_rows"] == 2
    assert data["successful_rows"] == 2
    assert data["failed_rows"] == 0
    assert len(data["students"]) == 2

    # API returns string UUIDs
    student_ids = {
        student["id"]
        for student in data["students"]
    }

    # Verify cohort memberships
    with Session(engine) as session:
        members = session.exec(
            select(CohortMember).where(
                CohortMember.cohort_id == cohort_id
            )
        ).all()

        assert len(members) == 2

        # Convert DB UUIDs to strings for comparison
        assigned_student_ids = {
            str(member.student_id)
            for member in members
        }

        assert assigned_student_ids == student_ids

        for member in members:
            assert member.org_id == org_id
            assert member.cohort_id == cohort_id
            assert member.added_by == user_id


# ============================================================
# STUDENT EXPORT
# ============================================================

EXPORT_URL = "/api/v1/users/students/export"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPECTED_EXPORT_HEADERS = [
    "ID", "First Name", "Last Name", "Registration Number",
    "Email", "Access Code", "Status", "Created At",
]


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _create_student(client: TestClient, token: str, **overrides: str) -> dict[str, Any]:
    """Create a student; overrides let tests vary phone / reg number / names."""
    payload = {**STUDENT_PAYLOAD, **overrides}
    response = client.post(
        "/api/v1/users/students/create", json=payload, headers=_auth(token),
    )
    assert response.status_code == 200, response.json()
    return response.json()


def _export(client: TestClient, token: str, **params: str) -> Response:
    return client.get(EXPORT_URL, params=params, headers=_auth(token))


def _read_export(response: Response) -> tuple[list[str], list[dict[str, Any]]]:
    """Parse the xlsx body into (headers, list of {header: value} records)."""
    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["Students"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:]]
    return headers, records


def _create_cohort(name: str = "EXPORT-COHORT") -> UUID:
    """Create a cohort in the signed-up super admin's org and return its id."""
    with Session(engine) as session:
        user = session.exec(
            select(UserModel).where(UserModel.email == USER_EMAIL)
        ).first()
        assert user is not None

        membership = session.exec(
            select(OrgMembership).where(OrgMembership.user_id == user.id)
        ).first()
        assert membership is not None

        cohort = CohortModel(
            org_id=membership.org_id,
            name=name,
            description="Export test cohort",
            created_by=user.id,
        )
        session.add(cohort)
        session.commit()
        session.refresh(cohort)
        return cohort.id


def test_admin_can_export_students(client: TestClient) -> None:
    token = get_super_admin_token(client)
    created = _create_student(client, token)

    response = _export(client, token)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX_MIME)
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"

    headers, records = _read_export(response)
    assert headers == EXPECTED_EXPORT_HEADERS
    assert len(records) == 1

    row = records[0]
    assert row["ID"] == created["id"]
    assert row["First Name"] == STUDENT_PAYLOAD["firstname"]
    assert row["Last Name"] == STUDENT_PAYLOAD["lastname"]
    assert row["Registration Number"] == STUDENT_PAYLOAD["institution_id"]
    assert row["Access Code"] == created["access_code"]
    assert row["Email"] is None  # students are created without an email
    assert isinstance(row["Status"], str) and row["Status"]
    assert isinstance(row["Created At"], datetime)


def test_export_contains_no_credentials(client: TestClient) -> None:
    """Regression: exports must never carry passwords or security Q&A."""
    token = get_super_admin_token(client)
    access_code = _create_student(client, token)["access_code"]
    client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": access_code,
            "favorite_question": "What is your pet's name?",
            "favorite_answer": "Fluffy",
        },
    )

    response = _export(client, token)
    assert response.status_code == 200, response.text

    headers, records = _read_export(response)
    lowered = [h.lower() for h in headers]
    assert not any(
        word in h for h in lowered for word in ("password", "question", "answer")
    )
    # The security answer must not leak into any cell either
    assert all("fluffy" not in str(v).lower() for r in records for v in r.values())


def test_export_does_not_modify_students(client: TestClient) -> None:
    """Export is read-only: a student can still log in afterwards."""
    token = get_super_admin_token(client)
    access_code = _create_student(client, token)["access_code"]
    setup = client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": access_code,
            "favorite_question": "What is your pet's name?",
            "favorite_answer": "Fluffy",
        },
    )
    assert setup.status_code == 200, setup.json()

    # reset_passwords was removed; a stale client sending it must change nothing
    export = _export(client, token, reset_passwords="true")
    assert export.status_code == 200, export.text

    login = client.post(
        "/api/v1/users/student/login",
        json={"access_code": access_code, "favorite_answer": "Fluffy"},
    )
    assert login.status_code == 200, login.json()


def test_export_filter_by_name(client: TestClient) -> None:
    token = get_super_admin_token(client)
    _create_student(client, token)  # Charlie Student
    _create_student(
        client, token,
        firstname="Diana", lastname="Okafor",
        phone="+1000000004", institution_id="STU/2024/002",
    )

    response = _export(client, token, name="diana")
    assert response.status_code == 200, response.text

    _, records = _read_export(response)
    assert len(records) == 1
    assert records[0]["First Name"] == "Diana"
    assert records[0]["Registration Number"] == "STU/2024/002"


def test_export_filter_by_cohort(client: TestClient) -> None:
    token = get_super_admin_token(client)
    cohort_id = _create_cohort()

    bulk = client.post(
        f"/api/v1/users/students/create/bulk?cohort_id={cohort_id}",
        files={"file": (
            "students.xlsx",
            create_student_excel([
                {"firstname": "In", "lastname": "CohortOne",
                 "email": "in1.export@example.com", "institution_id": "EXP/001"},
                {"firstname": "In", "lastname": "CohortTwo",
                 "email": "in2.export@example.com", "institution_id": "EXP/002"},
            ]).getvalue(),
            XLSX_MIME,
        )},
        headers=_auth(token),
    )
    assert bulk.status_code == 200, bulk.json()
    assert bulk.json()["successful_rows"] == 2

    _create_student(client, token)  # not in the cohort

    response = _export(client, token, cohort_id=str(cohort_id))
    assert response.status_code == 200, response.text

    _, records = _read_export(response)
    assert {r["Registration Number"] for r in records} == {"EXP/001", "EXP/002"}

    # Unfiltered export includes everyone
    _, all_records = _read_export(_export(client, token))
    assert len(all_records) == 3


def test_export_with_no_matches_returns_header_only(client: TestClient) -> None:
    token = get_super_admin_token(client)
    _create_student(client, token)

    response = _export(client, token, status="archived")
    assert response.status_code == 200, response.text

    headers, records = _read_export(response)
    assert headers == EXPECTED_EXPORT_HEADERS
    assert records == []


def test_export_neutralises_formula_injection(client: TestClient) -> None:
    """A name starting with '=' must not become a live Excel formula."""
    token = get_super_admin_token(client)
    _create_student(client, token, firstname="=1+1")

    response = _export(client, token)
    assert response.status_code == 200, response.text

    _, records = _read_export(response)
    assert records[0]["First Name"] == "'=1+1"


# def test_teacher_can_export_students(client: TestClient) -> None:
#     super_token = get_super_admin_token(client)
#     _create_student(client, super_token)

#     teacher_create = client.post(
#         "/api/v1/users/staff/create",
#         json=TEACHER_PAYLOAD,
#         headers=_auth(super_token),
#     )
#     assert teacher_create.status_code == 200, teacher_create.json()
#     teacher_token = _activate_staff(client, teacher_create.json()["id"])["access_token"]

#     response = _export(client, teacher_token)
#     assert response.status_code == 200, response.text
#     assert response.headers["content-type"].startswith(XLSX_MIME)
    # If teachers become cohort-scoped, assert the narrowed row count here.


def test_student_cannot_export_students(client: TestClient) -> None:
    super_token = get_super_admin_token(client)
    access_code = _create_student(client, super_token)["access_code"]

    setup = client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": access_code,
            "favorite_question": "What is your pet's name?",
            "favorite_answer": "Fluffy",
        },
    )
    assert setup.status_code == 200, setup.json()
    student_token = setup.json()["access_token"]

    response = _export(client, student_token)
    assert response.status_code == 403, response.text


def test_export_requires_authentication(client: TestClient) -> None:
    response = client.get(EXPORT_URL)
    assert response.status_code == 401


def test_old_activation_link_rejected_after_email_change(client):
    token = get_super_admin_token(client)
    teacher = _create_pending_teacher(client, token)

    teacher_id = teacher["id"]
    teacher_email = teacher["email"]

    # Token issued for the original email
    old_token = make_staff_activation_token(
        user_id=teacher_id,
        email=teacher_email,
    )

    patch_resp = client.patch(
        f"/api/v1/users/staff/{teacher_id}",
        json={"email": "corrected@cbtech.com"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert patch_resp.status_code == 200, patch_resp.json()

    payload = {
        "password": "newSecurePass123!",
        "confirm_password": "newSecurePass123!",
    }

    # Old token should now be rejected because it contains the old email.
    stale = client.post(
        "/api/v1/users/staff/activate",
        json={"token": old_token, **payload},
    )
    assert stale.status_code == 400, stale.json()

    # New token must contain the NEW email.
    fresh_token = make_staff_activation_token(
        user_id=teacher_id,
        email="corrected@cbtech.com",
    )

    fresh = client.post(
        "/api/v1/users/staff/activate",
        json={"token": fresh_token, **payload},
    )
    assert fresh.status_code == 200, fresh.json()


def test_admin_can_edit_pending_staff_email_and_resend(client):
    token = get_super_admin_token(client)
    teacher = _create_pending_teacher(client, token)
    teacher_id = teacher["id"]

    with patch(
        "auth.utility.email.email_service.EmailService.send_staff_activation_email",
        new_callable=AsyncMock,
    ) as send:
        resp = client.patch(
            f"/api/v1/users/staff/{teacher_id}/email",
            json={"email": "bob.new@cbtech.com"},
            headers={"Authorization": f"Bearer {token}"},
        )

    assert resp.status_code == 200, resp.json()
    data = resp.json()
    assert data["email"] == "bob.new@cbtech.com"
    assert data["email_changed"] is True
    assert data["email_sent"] is True
    send.assert_awaited_once()


def test_cannot_edit_email_of_activated_staff(client):
    token = get_super_admin_token(client)
    teacher = _create_pending_teacher(client, token)
    teacher_id = teacher["id"]
    teacher_email = teacher["email"]
    _activate_staff(client, teacher_id, teacher_email)

    resp = client.patch(
        f"/api/v1/users/staff/{teacher_id}/email",
        json={"email": "bob.new@cbtech.com"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 409, resp.json()










 # ============================================================
# STUDENT CREDENTIALS: ADMIN RESET + SELF-SERVICE Q&A CHANGE
# ============================================================

STUDENT_QUESTION = "What is your pet's name?"
STUDENT_ANSWER = "Fluffy"
TEMP_QUESTION = "Enter the temporary answer given to you by your administrator"
STUDENT_LOGIN_URL = "/api/v1/users/student/login"
STUDENT_QUESTION_URL = "/api/v1/users/student/login/question"
CHANGE_QA_URL = "/api/v1/users/students/me/security-qa"

VALID_QA_CHANGE: dict[str, str] = {
    "current_answer": STUDENT_ANSWER,
    "new_favorite_question": "What city were you born in?",
    "new_favorite_answer": "Lagos",
}


def _reset_url(student_id: str | UUID) -> str:
    return f"/api/v1/users/students/{student_id}/reset-credentials"


def _create_active_student(
    client: TestClient, token: str, **overrides: str
) -> dict[str, str]:
    """Create a student and complete first-time setup. Returns id, access_code, student token."""
    created = _create_student(client, token, **overrides)
    setup = client.post(
        "/api/v1/users/student/init",
        json={
            "access_code": created["access_code"],
            "favorite_question": STUDENT_QUESTION,
            "favorite_answer": STUDENT_ANSWER,
        },
    )
    assert setup.status_code == 200, setup.json()
    return {
        "id": created["id"],
        "access_code": created["access_code"],
        "token": setup.json()["access_token"],
    }


def _staff_token(client: TestClient, super_token: str, payload: dict[str, str]) -> str:
    created = client.post(
        "/api/v1/users/staff/create", json=payload, headers=_auth(super_token),
    )
    assert created.status_code == 200, created.json()
    return _activate_staff(client, created.json()["id"], payload["email"])["access_token"]


def _student_login(client: TestClient, access_code: str, answer: str) -> Response:
    return client.post(
        STUDENT_LOGIN_URL, json={"access_code": access_code, "favorite_answer": answer},
    )


def _get_question(client: TestClient, access_code: str) -> Response:
    return client.post(STUDENT_QUESTION_URL, json={"access_code": access_code})


def _reset(
    client: TestClient,
    token: str,
    student_id: str | UUID,
    body: dict[str, Any] | None = None,
) -> Response:
    if body is None:
        return client.post(_reset_url(student_id), headers=_auth(token))
    return client.post(_reset_url(student_id), json=body, headers=_auth(token))


def _change_qa(client: TestClient, token: str, **overrides: str) -> Response:
    return client.post(
        CHANGE_QA_URL, json={**VALID_QA_CHANGE, **overrides}, headers=_auth(token),
    )


# ------------------------------------------------------------
# Admin reset: success paths
# ------------------------------------------------------------

def test_admin_reset_generates_temp_answer(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, token, student["id"])
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["was_generated"] is True
    assert data["student_id"] == student["id"]
    assert data["access_code"] == student["access_code"]  # access code is untouched
    temp_answer: str = data["temporary_answer"]
    assert len(temp_answer) >= 6

    login = _student_login(client, student["access_code"], temp_answer)
    assert login.status_code == 200, login.text
    assert "access_token" in login.json()

    # The old answer is dead
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 401

    # The student now sees the placeholder question
    question = _get_question(client, student["access_code"])
    assert question.status_code == 200, question.text
    assert question.json()["favorite_question"] == TEMP_QUESTION


def test_admin_reset_with_supplied_temp_answer(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, token, student["id"], {"temp_answer": "blue-door-7"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["was_generated"] is False
    assert data["temporary_answer"] == "blue-door-7"

    assert _student_login(client, student["access_code"], "blue-door-7").status_code == 200
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 401


def test_supplied_temp_answer_is_trimmed(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, token, student["id"], {"temp_answer": "  blue-door-7  "})
    assert response.status_code == 200, response.text
    assert response.json()["temporary_answer"] == "blue-door-7"
    assert _student_login(client, student["access_code"], "blue-door-7").status_code == 200


def test_temp_answer_is_case_sensitive(client: TestClient) -> None:
    """Documents current behaviour: answers are compared exactly as typed."""
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)
    _reset(client, token, student["id"], {"temp_answer": "BlueDoor7"})

    assert _student_login(client, student["access_code"], "bluedoor7").status_code == 401
    assert _student_login(client, student["access_code"], "BlueDoor7").status_code == 200


@pytest.mark.parametrize(
    "body",
    [None, {}, {"temp_answer": None}, {"temp_answer": ""}, {"temp_answer": "   "}],
    ids=["no-body", "empty-object", "null", "empty-string", "whitespace"],
)
def test_blank_or_missing_temp_answer_is_generated(
    client: TestClient, body: dict[str, Any] | None
) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, token, student["id"], body)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["was_generated"] is True
    assert data["temporary_answer"]
    assert _student_login(
        client, student["access_code"], data["temporary_answer"]
    ).status_code == 200


def test_each_reset_generates_a_new_answer(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    first: str = _reset(client, token, student["id"]).json()["temporary_answer"]
    second: str = _reset(client, token, student["id"]).json()["temporary_answer"]

    assert first != second
    assert _student_login(client, student["access_code"], first).status_code == 401
    assert _student_login(client, student["access_code"], second).status_code == 200


def test_temp_answer_is_stored_hashed(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)
    _reset(client, token, student["id"], {"temp_answer": "blue-door-7"})

    with Session(engine) as session:
        user = session.get(UserModel, UUID(student["id"]))
        assert user is not None
        assert user.favorite_answer_hash
        assert user.favorite_answer_hash != "blue-door-7"
        assert PasswordHasher.verify("blue-door-7", user.favorite_answer_hash)
        assert user.favorite_question == TEMP_QUESTION
        assert user.is_first_login is False  # reset must not push them back into setup


# ------------------------------------------------------------
# Admin reset: validation and state errors
# ------------------------------------------------------------

def test_too_short_temp_answer_is_rejected(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, token, student["id"], {"temp_answer": "abc"})
    assert response.status_code == 422, response.text

    # Nothing changed
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 200


def test_reset_unknown_student_returns_404(client: TestClient) -> None:
    token = get_super_admin_token(client)
    response = _reset(client, token, uuid4())
    assert response.status_code == 404, response.text


def test_reset_rejects_non_student_target(client: TestClient) -> None:
    token = get_super_admin_token(client)
    teacher = _create_pending_teacher(client, token)

    response = _reset(client, token, teacher["id"])
    assert response.status_code == 404, response.text


def test_reset_before_first_time_setup_is_rejected(client: TestClient) -> None:
    token = get_super_admin_token(client)
    created = _create_student(client, token)  # no /student/init

    response = _reset(client, token, created["id"])
    assert response.status_code == 400, response.text
    assert "first-time setup" in response.json()["detail"]


def test_reset_rejects_inactive_membership(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    with Session(engine) as session:
        membership = session.exec(
            select(OrgMembership).where(OrgMembership.user_id == UUID(student["id"]))
        ).first()
        assert membership is not None
        membership.status = MembershipStatus.ARCHIVED
        session.add(membership)
        session.commit()

    response = _reset(client, token, student["id"])
    assert response.status_code == 400, response.text
    assert "first-time" not in response.json()["detail"]  # needs fix #1


def test_reset_is_scoped_to_the_admins_organization(client: TestClient) -> None:
    """A different org_id must not find the student (service-level: avoids a second signup)."""
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    with Session(engine) as session:
        with pytest.raises(HTTPException) as excinfo:
            UserManagementService.admin_reset_student_credentials(
                session, UUID(student["id"]), uuid4(),
            )
    assert excinfo.value.status_code == 404

    # Nothing changed
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 200


# ------------------------------------------------------------
# Admin reset: authorization
# ------------------------------------------------------------

def test_plain_admin_can_reset_student_credentials(client: TestClient) -> None:
    super_token = get_super_admin_token(client)
    student = _create_active_student(client, super_token)
    admin_token = _staff_token(client, super_token, ADMIN_PAYLOAD)

    response = _reset(client, admin_token, student["id"])
    assert response.status_code == 200, response.text


def test_teacher_cannot_reset_student_credentials(client: TestClient) -> None:
    super_token = get_super_admin_token(client)
    student = _create_active_student(client, super_token)
    teacher_token = _staff_token(client, super_token, TEACHER_PAYLOAD)

    response = _reset(client, teacher_token, student["id"])
    assert response.status_code == 403, response.text
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 200


def test_student_cannot_reset_credentials(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _reset(client, student["token"], student["id"])
    assert response.status_code == 403, response.text


def test_reset_requires_authentication(client: TestClient) -> None:
    response = client.post(_reset_url(uuid4()))
    assert response.status_code == 401, response.text


# ------------------------------------------------------------
# Student self-service Q&A change
# ------------------------------------------------------------

def test_student_can_change_own_security_qa(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _change_qa(client, student["token"])
    assert response.status_code == 200, response.text
    assert response.json() == {"detail": "Security question and answer updated."}

    assert _student_login(client, student["access_code"], "Lagos").status_code == 200
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 401
    question = _get_question(client, student["access_code"])
    assert question.json()["favorite_question"] == "What city were you born in?"


def test_change_qa_with_wrong_current_answer_is_rejected(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _change_qa(client, student["token"], current_answer="Wrong")
    assert response.status_code == 400, response.text  # 400, not 401: must not log the student out
    assert response.json()["detail"] == "Current answer is incorrect."

    # Nothing changed
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 200
    assert _get_question(client, student["access_code"]).json()["favorite_question"] == STUDENT_QUESTION


def test_change_qa_trims_the_question(client: TestClient) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _change_qa(
        client, student["token"], new_favorite_question="  What city were you born in?  ",
    )
    assert response.status_code == 200, response.text
    question = _get_question(client, student["access_code"])
    assert question.json()["favorite_question"] == "What city were you born in?"


@pytest.mark.parametrize(
    "payload",
    [
        {"new_favorite_question": "What city?", "new_favorite_answer": "Lagos"},
        {"current_answer": "", "new_favorite_question": "What city?", "new_favorite_answer": "Lagos"},
        {"current_answer": STUDENT_ANSWER, "new_favorite_answer": "Lagos"},
        {"current_answer": STUDENT_ANSWER, "new_favorite_question": "ab", "new_favorite_answer": "Lagos"},
        {"current_answer": STUDENT_ANSWER, "new_favorite_question": "What city?", "new_favorite_answer": "a"},
        {"current_answer": STUDENT_ANSWER, "new_favorite_question": "What city?", "new_favorite_answer": "x" * 101},
    ],
    ids=[
        "missing-current", "empty-current", "missing-question",
        "question-too-short", "answer-too-short", "answer-too-long",
    ],
)
def test_change_qa_rejects_invalid_payloads(
    client: TestClient, payload: dict[str, str]
) -> None:
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = client.post(CHANGE_QA_URL, json=payload, headers=_auth(student["token"]))
    assert response.status_code == 422, response.text
    assert _student_login(client, student["access_code"], STUDENT_ANSWER).status_code == 200


def test_change_qa_rejects_blank_question(client: TestClient) -> None:
    """Needs fix #2: '   ' must not pass min_length and then be stored as ''."""
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)

    response = _change_qa(client, student["token"], new_favorite_question="   ")
    assert response.status_code == 422, response.text
    assert _get_question(client, student["access_code"]).json()["favorite_question"] == STUDENT_QUESTION


def test_student_can_change_qa_after_admin_reset(client: TestClient) -> None:
    """End-to-end: forgot answer -> admin reset -> student logs in -> sets their own."""
    token = get_super_admin_token(client)
    student = _create_active_student(client, token)
    _reset(client, token, student["id"], {"temp_answer": "blue-door-7"})

    login = _student_login(client, student["access_code"], "blue-door-7")
    assert login.status_code == 200, login.text
    fresh_token: str = login.json()["access_token"]

    response = _change_qa(client, fresh_token, current_answer="blue-door-7")
    assert response.status_code == 200, response.text

    assert _student_login(client, student["access_code"], "Lagos").status_code == 200
    assert _student_login(client, student["access_code"], "blue-door-7").status_code == 401


def test_change_qa_does_not_affect_other_students(client: TestClient) -> None:
    token = get_super_admin_token(client)
    first = _create_active_student(client, token)
    second = _create_active_student(
        client, token,
        firstname="Diana", lastname="Okafor",
        phone="+1000000004", institution_id="STU/2024/002",
    )

    assert _change_qa(client, first["token"]).status_code == 200

    # The other student's credentials are untouched
    assert _student_login(client, second["access_code"], STUDENT_ANSWER).status_code == 200
    assert _get_question(client, second["access_code"]).json()["favorite_question"] == STUDENT_QUESTION


def test_teacher_cannot_change_student_qa(client: TestClient) -> None:
    super_token = get_super_admin_token(client)
    teacher_token = _staff_token(client, super_token, TEACHER_PAYLOAD)

    response = _change_qa(client, teacher_token)
    assert response.status_code == 403, response.text


def test_change_qa_requires_authentication(client: TestClient) -> None:
    response = client.post(CHANGE_QA_URL, json=VALID_QA_CHANGE)
    assert response.status_code == 401, response.text